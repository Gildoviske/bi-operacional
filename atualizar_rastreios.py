# -*- coding: utf-8 -*-
"""Lê os relatórios de vendas dos Correios (CLI*.TXT) e alimenta a aba
"RASTREIO DOS CORREIOS" da planilha TRANSFERÊNCIAS PENDENTES.xlsx.

Cada postagem é classificada pelo CEP de destino: se o CEP é de uma filial, o envio
foi para a rede; se não é, foi para cliente. O CEP das filiais vem do diretório
(diretorio-tim/gerar_diretorio.py) e o nome, da planilha de contatos dos gerentes.

A gravação é feita direto no XML da aba, sem passar pelo openpyxl, porque a planilha
tem um suplemento do Office (web extension) que o openpyxl descarta ao salvar.
"""
import io
import json
import re
import shutil
import unicodedata
import sys
import urllib.request
import zipfile
from datetime import date, datetime
from pathlib import Path

import openpyxl

import correios_api

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

BASE = Path(r"C:\Users\recru\Desktop\CONTROLE COMPRAS")
# a pasta de rastreios fica no Google Drive: é por lá que chegam os relatórios dos Correios
PASTA_TXT = Path(r"I:\Meu Drive\CONTROLE COMPRAS\rastreios")
PLANILHA = BASE / "TRANSFERÊNCIAS PENDENTES.xlsx"
ABA_XML = "xl/worksheets/sheet2.xml"          # RASTREIO DOS CORREIOS
DIRETORIO = Path(r"C:\Users\recru\Desktop\diretorio-tim")
CONTATOS = Path(r"I:\Meu Drive\CONTATOS DOS GERENTES - LOJAS TIM.xlsx")
CACHE = PASTA_TXT / "cache_ceps.json"

ESTILO_DATA = "8"       # dd/mm/aaaa, o mesmo usado na aba de transferências
ESTILO_CABECALHO = "10"

COLUNAS = [
    ("DATA", 12), ("O.S.", 10), ("PRODUTO", 20), ("RASTREIO", 15), ("CEP", 11),
    ("TIPO", 10), ("FILIAL", 30), ("DESTINATÁRIO", 30), ("CIDADE/UF", 22),
    ("ENDEREÇO", 38), ("VALOR", 10), ("ARQUIVO", 14),
    ("STATUS ENTREGA", 34), ("LOCAL ATUAL", 22), ("ATUALIZADO EM", 17),
    ("PREVISÃO ENTREGA", 18),
]

# uma postagem por linha: ...valor  RASTREIO  CEP  destinatário. O rastreio e o CEP
# ancoram a leitura porque data e nº da O.S. só aparecem na 1ª linha de cada grupo.
LINHA = re.compile(
    r"^(?P<cabeca>.*?)"
    r"(?P<valor>\d{1,3}(?:\.\d{3})*,\d{2})\s+"
    r"(?P<rastreio>[A-Z]{2}\d{9})\s+"
    r"(?P<cep>\d{1,5}-\d{3})\s*"
    r"(?P<destinatario>.*?)\s*$"
)
DATA_LINHA = re.compile(r"^\s*(\d{1,2})/\s*(\d{1,2})")
OS_PRODUTO = re.compile(r"^\s*(?P<os>\d+)?\s+(?P<produto>[A-Z][A-Z0-9 /]*?)\s{2,}")
PERIODO = re.compile(r"Per[ií]odo:\s*(\d{2}/\d{2}/\d{4})\s*a\s*(\d{2}/\d{2}/\d{4})")
CLIENTE_TXT = re.compile(r"^Cliente\s+(.+?)\s{2,}CNPJ", re.MULTILINE)

# postagem endereçada ao próprio grupo é remessa interna, não venda para cliente.
# o titular da conta sai do cabeçalho de cada TXT; aqui ficam as demais razões sociais.
RAZOES_PROPRIAS = {"F M M ROCHA", "FMM ROCHA", "ROCHA TELECOM"}

# uma filial pode receber em mais de um CEP: o cadastrado no diretorio e outros que
# aparecem nas remessas. CEP extra -> codigo SAP da filial.
CEPS_EXTRAS = {
    "13271600": "44924",   # Valinhos, usado nas remessas de contrato
    "09390040": "ADM",     # Outlet do Celular, mesma avenida do cadastro (09390-120)
}


def normalizar_nome(texto):
    s = unicodedata.normalize("NFKD", str(texto or "").upper()).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"[^A-Z0-9 ]", " ", s).replace(" LTDA", "").replace(" EPP", "")


def e_razao_propria(destinatario, razoes):
    nome = " ".join(normalizar_nome(destinatario).split())
    return any(r and r in nome for r in razoes)


def filiais_por_cep():
    """CEP (8 dígitos) -> {sap, filial, endereco}."""
    fonte = io.open(DIRETORIO / "gerar_diretorio.py", encoding="utf-8").read()
    bloco = fonte[fonte.index("ENDERECOS_MANUAIS = {"):]
    bloco = bloco[:bloco.index("\n}")]
    enderecos = re.findall(
        r'"([A-Za-z0-9]+)":\s*\{"endereco":\s*"(.*?)",\s*"cep":\s*"([\d-]+)"\}', bloco)

    # o SAP vem da planilha ora como texto, ora como numero ("23173" ou 23173.0)
    def so_digitos(valor):
        s = str(valor).strip()
        return s[:-2] if s.endswith(".0") else s

    wb = openpyxl.load_workbook(CONTATOS, data_only=True, read_only=True)
    nomes = {}
    for linha in wb[wb.sheetnames[0]].iter_rows(values_only=True):
        campos = (list(linha) + [None] * 8)[:8]
        filial, sap = campos[2], campos[6]
        if filial and sap:
            nomes[so_digitos(sap)] = str(filial).replace("ROCHA TELECOM - ", "").strip()

    mapa = {
        cep.replace("-", ""): {
            "sap": sap, "filial": nomes.get(sap, f"SAP {sap}"), "endereco": endereco,
        }
        for sap, endereco, cep in enderecos
    }
    por_sap = {d["sap"]: d for d in mapa.values()}
    for cep_extra, sap in CEPS_EXTRAS.items():
        if sap in por_sap and cep_extra not in mapa:
            mapa[cep_extra] = por_sap[sap]

    sem_nome = sorted({d["sap"] for d in mapa.values() if d["filial"].startswith("SAP ")})
    if sem_nome:
        print(f"   aviso: SAP sem nome na planilha de contatos: {', '.join(sem_nome)}")
    return mapa


def ano_da_postagem(dia, mes, inicio, fim):
    """O TXT traz só dia/mês; o ano vem do período do cabeçalho."""
    for ano in {inicio.year, fim.year}:
        try:
            d = date(ano, mes, dia)
        except ValueError:
            continue
        if inicio <= d <= fim:
            return d
    try:
        return date(inicio.year, mes, dia)
    except ValueError:
        return None


def ler_txt(caminho):
    texto = caminho.read_bytes().decode("cp1252")
    linhas = texto.splitlines()

    m = PERIODO.search(texto)
    if m:
        inicio = datetime.strptime(m.group(1), "%d/%m/%Y").date()
        fim = datetime.strptime(m.group(2), "%d/%m/%Y").date()
    else:
        hoje = date.today()
        inicio = fim = hoje
        print(f"   aviso: {caminho.name} sem linha de Período, usando {hoje}")

    titular = CLIENTE_TXT.search(texto)
    if titular:
        RAZOES_PROPRIAS.add(" ".join(normalizar_nome(titular.group(1)).split()))

    itens = []
    data_atual = os_atual = None
    for linha in linhas:
        m = LINHA.match(linha)
        if not m:
            continue
        cabeca = m.group("cabeca")
        md = DATA_LINHA.match(cabeca)
        if md:
            data_atual = ano_da_postagem(int(md.group(1)), int(md.group(2)), inicio, fim)
            cabeca = cabeca[md.end():]
        mo = OS_PRODUTO.match(cabeca)
        if mo and mo.group("os"):
            os_atual = mo.group("os")
        produto = mo.group("produto").strip() if mo else cabeca.strip()
        itens.append({
            "data": data_atual,
            "os": os_atual or "-",
            "produto": produto,
            # o TXT traz o codigo sem o sufixo de pais; com BR fica igual ao dos Correios
            "rastreio": m.group("rastreio") + "BR",
            "cep": m.group("cep").replace("-", "").rjust(8, "0"),
            "destinatario": m.group("destinatario").strip() or "-",
            "valor": float(m.group("valor").replace(".", "").replace(",", ".")),
            "arquivo": caminho.name,
        })
    return itens


def buscar_endereco(cep, cache):
    """CEP pela API oficial dos Correios; ViaCEP fica como reserva."""
    if cep in cache:
        return cache[cep]
    oficial = correios_api.buscar_cep(cep)
    if oficial:
        cache[cep] = oficial
        return oficial
    try:
        req = urllib.request.Request(
            f"https://viacep.com.br/ws/{cep}/json/", headers={"User-Agent": "Mozilla/5.0"})
        with urllib.request.urlopen(req, timeout=8) as resp:
            dados = json.loads(resp.read().decode("utf-8"))
        endereco = None if dados.get("erro") else {
            "logradouro": dados.get("logradouro") or "", "bairro": dados.get("bairro") or "",
            "cidade": dados.get("localidade") or "", "uf": dados.get("uf") or "",
        }
    except Exception:
        return None
    cache[cep] = endereco
    return endereco


def escapar(texto):
    return (str(texto).replace("&", "&amp;").replace("<", "&lt;")
            .replace(">", "&gt;").replace('"', "&quot;"))


def col_letra(n):
    letra = ""
    while n > 0:
        n, resto = divmod(n - 1, 26)
        letra = chr(65 + resto) + letra
    return letra


def celula(ref, valor, estilo=None, numero=False):
    s = f' s="{estilo}"' if estilo else ""
    if numero:
        return f'<c r="{ref}"{s}><v>{valor}</v></c>'
    return f'<c r="{ref}"{s} t="inlineStr"><is><t xml:space="preserve">{escapar(valor)}</t></is></c>'


def montar_sheet_xml(itens):
    """Gera o XML da aba do zero, com texto inline para não mexer no sharedStrings."""
    epoca = date(1899, 12, 30)
    linhas_xml = []

    cabecalho = "".join(
        celula(f"{col_letra(i + 1)}1", titulo, ESTILO_CABECALHO)
        for i, (titulo, _) in enumerate(COLUNAS))
    linhas_xml.append(f'<row r="1">{cabecalho}</row>')

    for n, it in enumerate(itens, start=2):
        campos = []
        if it["data"]:
            campos.append(celula(f"A{n}", (it["data"] - epoca).days, ESTILO_DATA, numero=True))
        else:
            campos.append(celula(f"A{n}", "-"))
        campos.append(celula(f"B{n}", it["os"]))
        campos.append(celula(f"C{n}", it["produto"]))
        campos.append(celula(f"D{n}", it["rastreio"]))
        campos.append(celula(f"E{n}", f'{it["cep"][:5]}-{it["cep"][5:]}'))
        campos.append(celula(f"F{n}", it["tipo"]))
        campos.append(celula(f"G{n}", it["filial"]))
        campos.append(celula(f"H{n}", it["destinatario"]))
        campos.append(celula(f"I{n}", it["cidade_uf"]))
        campos.append(celula(f"J{n}", it["endereco"]))
        campos.append(celula(f"K{n}", f'{it["valor"]:.2f}', numero=True))
        campos.append(celula(f"L{n}", it["arquivo"]))
        campos.append(celula(f"M{n}", it.get("status_entrega") or "-"))
        campos.append(celula(f"N{n}", it.get("local_atual") or "-"))
        campos.append(celula(f"O{n}", it.get("rastreado_em") or "-"))
        prev = it.get("previsao_data")
        if prev:
            campos.append(celula(f"P{n}", (prev - epoca).days, ESTILO_DATA, numero=True))
        else:
            campos.append(celula(f"P{n}", "-"))
        linhas_xml.append(f'<row r="{n}">{"".join(campos)}</row>')

    ultima = len(itens) + 1
    cols = "".join(
        f'<col min="{i + 1}" max="{i + 1}" width="{largura}" customWidth="1"/>'
        for i, (_, largura) in enumerate(COLUNAS))

    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
        f'<dimension ref="A1:{col_letra(len(COLUNAS))}{ultima}"/>'
        '<sheetViews><sheetView workbookViewId="0">'
        '<pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>'
        '</sheetView></sheetViews>'
        '<sheetFormatPr defaultRowHeight="15"/>'
        f'<cols>{cols}</cols>'
        f'<sheetData>{"".join(linhas_xml)}</sheetData>'
        f'<autoFilter ref="A1:{col_letra(len(COLUNAS))}{ultima}"/>'
        '<pageMargins left="0.7" right="0.7" top="0.75" bottom="0.75" header="0.3" footer="0.3"/>'
        '</worksheet>'
    )


def gravar_aba(xml_novo):
    """Reescreve o .xlsx trocando só o XML da aba, preservando todo o resto."""
    carimbo = datetime.now().strftime("%Y-%m-%d %H%M")
    backup = PLANILHA.with_name(f"{PLANILHA.stem} (backup {carimbo}).xlsx")
    shutil.copy2(PLANILHA, backup)

    origem = zipfile.ZipFile(PLANILHA, "r")
    partes = [(i, origem.read(i.filename)) for i in origem.infolist()]
    origem.close()
    assert any(i.filename == ABA_XML for i, _ in partes), f"{ABA_XML} não existe na planilha"

    temporario = PLANILHA.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(temporario, "w", zipfile.ZIP_DEFLATED) as destino:
        for info, dados in partes:
            if info.filename == ABA_XML:
                dados = xml_novo.encode("utf-8")
            destino.writestr(info, dados)
    temporario.replace(PLANILHA)
    return backup


def main():
    # no Windows o glob não diferencia maiúsculas, então *.TXT e *.txt trariam o
    # mesmo arquivo duas vezes
    arquivos = sorted({p.resolve() for p in PASTA_TXT.iterdir()
                       if p.is_file() and p.suffix.lower() == ".txt"})
    if not arquivos:
        raise SystemExit(f"nenhum .TXT em {PASTA_TXT}")

    itens = []
    for arquivo in arquivos:
        lidos = ler_txt(arquivo)
        print(f"   {arquivo.name}: {len(lidos)} postagens")
        itens.extend(lidos)

    # o mesmo rastreio pode vir em arquivos de períodos sobrepostos
    unicos = {}
    for it in itens:
        unicos.setdefault(it["rastreio"], it)
    repetidas = len(itens) - len(unicos)
    itens = sorted(unicos.values(), key=lambda i: (i["data"] or date.min, i["rastreio"]), reverse=True)
    print(f"   {len(itens)} postagens únicas em {len(arquivos)} arquivo(s)"
          + (f", {repetidas} repetidas descartadas" if repetidas else ""))

    por_cep = filiais_por_cep()
    cache = correios_api.ler_json_seguro(CACHE, {})
    # o diretório já consultou os CEPs das filiais; aproveita como semente
    semente = DIRETORIO / "cache_ceps.json"
    if semente.exists():
        for k, v in json.loads(semente.read_text(encoding="utf-8")).items():
            cache.setdefault(k, v)

    # "Av. Guarda Mor Lobo Viana, 427 - Centro, São Sebastião - SP" -> "São Sebastião/SP"
    cidades_de_filial = {}
    for dados in por_cep.values():
        final = dados["endereco"].split(",")[-1].strip()
        dados["cidade_uf"] = final.replace(" - ", "/")
        cidades_de_filial.setdefault(final.split(" - ")[0].strip().upper(), []).append(dados["filial"])

    consultas = 0
    suspeitos = []
    internos = []
    for it in itens:
        filial = por_cep.get(it["cep"])
        interno = e_razao_propria(it["destinatario"], RAZOES_PROPRIAS)
        if filial:
            it["tipo"] = "FILIAL"
            it["filial"] = filial["filial"]
            it["endereco"] = filial["endereco"]
            it["cidade_uf"] = filial["cidade_uf"]
        elif interno:
            # endereçado ao próprio grupo: é remessa interna mesmo sem o CEP bater.
            # tenta identificar a loja pela cidade, quando ela só tem uma.
            it["tipo"] = "FILIAL"
            end = buscar_endereco(it["cep"], cache)
            cidade = (end or {}).get("cidade", "")
            candidatas = cidades_de_filial.get(cidade.strip().upper(), [])
            it["filial"] = candidatas[0] if len(candidatas) == 1 else it["destinatario"]
            it["cidade_uf"] = f'{cidade}/{end["uf"]}' if end else "-"
            it["endereco"] = ", ".join(p for p in ((end or {}).get("logradouro"), (end or {}).get("bairro")) if p) or "-"
            internos.append((it["rastreio"], it["cep"], it["filial"], cidade))
        else:
            it["tipo"] = "CLIENTE"
            it["filial"] = "-"
            antes = it["cep"] in cache
            end = buscar_endereco(it["cep"], cache)
            consultas += 0 if antes else 1
            if end:
                it["cidade_uf"] = f'{end["cidade"]}/{end["uf"]}'
                it["endereco"] = ", ".join(p for p in (end["logradouro"], end["bairro"]) if p)
                # cidade que só tem uma loja provavelmente é envio para ela com o CEP
                # divergente do cadastro — não reclassifica sozinho, só avisa
                candidatas = cidades_de_filial.get(end["cidade"].strip().upper(), [])
                if len(candidatas) == 1:
                    suspeitos.append((it["rastreio"], it["cep"], candidatas[0], end["cidade"]))
            else:
                it["cidade_uf"] = it["endereco"] = "-"

    correios_api.gravar_json_seguro(CACHE, cache)

    # rastreamento oficial: sem credencial preenchida, segue com o que houver em cache
    rastreamento = correios_api.atualizar([i["rastreio"] for i in itens])
    for it in itens:
        info = rastreamento.get(it["rastreio"]) or {}
        it["status_entrega"] = info.get("status") or "-"
        it["local_atual"] = info.get("local") or "-"
        it["rastreado_em"] = info.get("data") or "-"
        it["entregue"] = bool(info.get("entregue"))
        try:
            it["previsao_data"] = date.fromisoformat(info["previsao"]) if info.get("previsao") else None
        except Exception:
            it["previsao_data"] = None

    backup = gravar_aba(montar_sheet_xml(itens))

    filiais = [i for i in itens if i["tipo"] == "FILIAL"]
    clientes = [i for i in itens if i["tipo"] == "CLIENTE"]
    print()
    if internos:
        print("Remessas internas (destinatario e do proprio grupo) com CEP fora do cadastro:")
        for rastreio, cep, filial, cidade in internos:
            print(f"   {rastreio}  {cep[:5]}-{cep[5:]}  {cidade} -> {filial}")
        print()
    if suspeitos:
        print("ATENCAO: postagens classificadas como CLIENTE numa cidade que so tem uma loja.")
        print("Se for envio para a filial, corrija o CEP dela em diretorio-tim/gerar_diretorio.py:")
        for rastreio, cep, filial, cidade in suspeitos:
            print(f"   {rastreio}  {cep[:5]}-{cep[5:]}  {cidade} -> unica loja na cidade: {filial}")
        print()
    print(f"backup: {backup.name}")
    entregues = sum(1 for i in itens if i.get("entregue"))
    print(f"OK: {len(itens)} postagens na aba RASTREIO DOS CORREIOS "
          f"({len(filiais)} para filiais, {len(clientes)} para clientes, "
          f"R$ {sum(i['valor'] for i in itens):,.2f} em frete, {consultas} CEPs consultados, "
          f"{entregues} entregues)")


if __name__ == "__main__":
    main()
